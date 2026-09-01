#!/usr/bin/env python3
"""
AI 聊天记录导出工具
支持从 DeepSeek 分享链接提取对话，导出为 Markdown 或 Word 文档。

用法:
    python extract_chat.py <分享链接> [-o 输出路径] [-f 格式] [--headed]

支持平台: DeepSeek (chat.deepseek.com/share/...)
"""

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from playwright.sync_api import sync_playwright


def sanitize_filename(name, max_len=60):
    """清理文件名中的非法字符。"""
    # 替换 Windows 文件名非法字符
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name)
    # 替换连续空格/下划线
    name = re.sub(r'[_\s]+', '_', name)
    # 去除首尾下划线和空格
    name = name.strip('_ ')
    # 截断过长名称
    if len(name) > max_len:
        name = name[:max_len].rstrip('_ ')
    return name


def extract_deepseek(page, url, timeout=30000):
    """提取 DeepSeek 分享链接中的对话内容及元信息。"""
    page.goto(url, wait_until="networkidle", timeout=timeout)

    # 等待对话内容加载
    try:
        page.wait_for_selector(".ds-message", timeout=timeout)
    except Exception:
        pass

    page.wait_for_timeout(3000)

    # 提取元信息
    meta = page.evaluate("""() => {
        const info = {};
        
        // 页面标题
        info.pageTitle = document.title || '';
        
        // 第一条用户消息（作为对话主题）
        const firstUserMsg = document.querySelector('.ds-collapsible-text, .fbb737a4');
        info.firstMessage = firstUserMsg ? firstUserMsg.innerText.substring(0, 150) : '';
        
        // og 标签
        const ogTitle = document.querySelector('meta[property="og:title"]');
        info.ogTitle = ogTitle ? ogTitle.content : '';
        
        // 文件附件名称
        const fileEl = document.querySelector('._967f3f9, .e70accd6');
        info.fileName = fileEl ? fileEl.innerText.trim() : '';
        
        // 从思考过程中提取论文标题（常见模式：包含 "Predictive"、"design" 等学术词汇）
        const thinkEls = document.querySelectorAll('.ds-think-content');
        let paperTitle = '';
        for (const el of thinkEls) {
            const text = el.innerText || '';
            // 尝试匹配论文标题模式（通常在 "题目：" 或引号中）
            const titleMatch = text.match(/题目[：:](.+)/);
            if (titleMatch) {
                paperTitle = titleMatch[1].trim();
                break;
            }
            // 匹配常见学术标题模式（英文）
            const engMatch = text.match(/(Predictive design.+?)(?:[.。\\n]|$)/i);
            if (engMatch) {
                paperTitle = engMatch[1].trim();
                break;
            }
        }
        info.paperTitle = paperTitle;
        
        return info;
    }""")

    # 使用 JavaScript 在页面内提取所有消息
    messages = page.evaluate("""() => {
        // 深度遍历节点，将公式块整体提取为 LaTeX 文本
        function extractNodeText(node) {
            if (node.nodeType === 3) return node.textContent;
            if (node.nodeType !== 1) return '';

            const tag = node.tagName;
            const cls = (typeof node.className === 'string') ? node.className : '';

            if (['SCRIPT', 'STYLE', 'BUTTON', 'SVG', 'PATH'].includes(tag)) return '';

            // 公式容器
            if (cls.includes('katex') || cls.includes('math') || cls.includes('equation') || tag === 'MATH') {
                const ann = node.querySelector('annotation[encoding="application/x-tex"]');
                if (ann) return '$' + ann.textContent + '$';
                return '$' + node.innerText.replace(/\\s+/g, ' ').trim() + '$';
            }

            // 代码块
            if (tag === 'CODE' || tag === 'PRE') {
                return '`' + node.innerText + '`';
            }

            // 表格
            if (tag === 'TABLE') {
                return extractTable(node);
            }

            // 图片
            if (tag === 'IMG') {
                return '![' + (node.alt || '') + '](' + (node.src || '') + ')';
            }

            if (tag === 'BR') return '\\n';

            const blockTags = ['P', 'DIV', 'H1', 'H2', 'H3', 'H4', 'H5', 'H6',
                               'LI', 'TR', 'BLOCKQUOTE', 'SECTION', 'ARTICLE'];
            const isBlock = blockTags.includes(tag);

            let result = '';
            for (const child of node.childNodes) {
                result += extractNodeText(child);
            }

            if (!cls.includes('katex') && !cls.includes('math')) {
                result = result.replace(/  +/g, ' ');
            }

            if (isBlock && result.trim()) {
                result = '\\n' + result.trim() + '\\n';
            }

            return result;
        }

        function extractTable(table) {
            const rows = [];
            for (const tr of table.querySelectorAll('tr')) {
                const cells = [];
                for (const td of tr.querySelectorAll('td, th')) {
                    cells.push(td.innerText.replace(/\\|/g, '\\\\|').trim());
                }
                rows.push(cells);
            }
            if (rows.length === 0) return '';
            let md = '| ' + rows[0].join(' | ') + ' |\\n';
            md += '| ' + rows[0].map(() => '---').join(' | ') + ' |\\n';
            for (let i = 1; i < rows.length; i++) {
                md += '| ' + rows[i].join(' | ') + ' |\\n';
            }
            return '\\n' + md;
        }

        const results = [];
        const allMsgs = document.querySelectorAll('.ds-message');
        for (const msg of allMsgs) {
            const aiMain = msg.querySelector('.ds-assistant-message-main-content');
            if (aiMain) {
                let mainContent = extractNodeText(aiMain);

                const thinkEl = msg.querySelector('.ds-think-content');
                if (thinkEl) {
                    const thinkText = extractNodeText(thinkEl);
                    if (mainContent.startsWith(thinkText)) {
                        mainContent = mainContent.substring(thinkText.length).trim();
                    }
                }

                mainContent = mainContent.replace(/\\n{3,}/g, '\\n\\n').trim();

                if (mainContent && mainContent.length > 5) {
                    results.push({ role: 'assistant', content: mainContent });
                }
            } else {
                const userEl = msg.querySelector('.ds-collapsible-text') ||
                               msg.querySelector('.fbb737a4') ||
                               msg.querySelector('._25c7358');
                if (userEl) {
                    const content = userEl.innerText.trim();
                    if (content && content.length > 2) {
                        results.push({ role: 'user', content: content });
                    }
                }
            }
        }
        return results;
    }""")

    return messages, meta


def parse_messages_to_markdown(messages, title="DeepSeek 对话记录"):
    """将消息列表转换为 Markdown 格式，一问一答配对。
    支持将一条包含多个问题的消息拆分为多个独立问答对。
    """
    import re as _re

    def split_questions(text):
        """将包含多个编号问题的文本拆分为独立问题列表。"""
        parts = _re.split(r'\n(?=\d+[、.．]\s)', text.strip())
        if len(parts) <= 1:
            parts = _re.split(r'\n(?=(?:问题|Q)\s*\d+)', text.strip())
        return [p.strip() for p in parts if p.strip()]

    def split_answers_by_questions(text, expected_count):
        """根据问题数量拆分答案。只在答案段落数与问题数匹配时才拆分。"""
        parts = _re.split(r'\n(?=\d+[.．：:]\s)', text.strip())
        if len(parts) <= 1:
            parts = _re.split(r'\n(?=\d+\s)', text.strip())
        if len(parts) == expected_count:
            return [p.strip() for p in parts if p.strip()]
        return [text.strip()]

    md_lines = [f"# {title}", ""]

    i = 0
    q_num = 0
    while i < len(messages):
        msg = messages[i]
        if msg["role"] == "user":
            questions = split_questions(msg["content"])
            if i + 1 < len(messages) and messages[i + 1]["role"] == "assistant":
                answers = split_answers_by_questions(messages[i + 1]["content"], len(questions))
                pairs = max(len(questions), len(answers))
                for j in range(pairs):
                    q_num += 1
                    q_text = questions[j] if j < len(questions) else "（见上文）"
                    a_text = answers[j] if j < len(answers) else "（见上文）"
                    md_lines.append(f"## 问题 {q_num}")
                    md_lines.append("")
                    md_lines.append(q_text)
                    md_lines.append("")
                    md_lines.append(f"## 回答 {q_num}")
                    md_lines.append("")
                    md_lines.append(a_text)
                    md_lines.append("")
                    md_lines.append("---")
                    md_lines.append("")
                i += 2
            else:
                for q_text in questions:
                    q_num += 1
                    md_lines.append(f"## 问题 {q_num}")
                    md_lines.append("")
                    md_lines.append(q_text)
                    md_lines.append("")
                    md_lines.append("---")
                    md_lines.append("")
                i += 1
        else:
            md_lines.append("## 回答")
            md_lines.append("")
            md_lines.append(msg["content"])
            md_lines.append("")
            md_lines.append("---")
            md_lines.append("")
            i += 1

    return "\n".join(md_lines)


def markdown_to_docx(md_text, output_path):
    """将 Markdown 转换为 Word 文档。"""
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.oxml.ns import qn
    except ImportError:
        print("提示: 需要安装 python-docx 才能导出 Word 文档")
        print("运行: pip install python-docx")
        return False

    doc = Document()

    # 设置默认字体为 Times New Roman
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(11)
    # 设置中文字体（东亚字体）
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), '宋体')

    def set_run_font(run, font_name="Times New Roman", east_asia="宋体", size=None):
        """为 run 设置字体。"""
        run.font.name = font_name
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = rPr.makeelement(qn('w:rFonts'), {})
            rPr.insert(0, rFonts)
        rFonts.set(qn('w:eastAsia'), east_asia)
        if size:
            run.font.size = Pt(size)

    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith("# "):
            heading = doc.add_heading(level=1)
            run = heading.add_run(line[2:])
            set_run_font(run, size=18)
        elif line.startswith("## "):
            heading = doc.add_heading(level=2)
            run = heading.add_run(line[3:])
            set_run_font(run, size=14)
        elif line == "---":
            doc.add_paragraph("")
        elif line.strip() == "":
            pass
        else:
            # 检查是否是代码块
            if line.strip().startswith("```"):
                code_lines = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    code_lines.append(lines[i])
                    i += 1
                if code_lines:
                    p = doc.add_paragraph()
                    run = p.add_run("\n".join(code_lines))
                    set_run_font(run, font_name="Consolas", east_asia="Consolas", size=9)
                    p.paragraph_format.left_indent = Inches(0.5)
            else:
                p = doc.add_paragraph()
                run = p.add_run(line)
                set_run_font(run)

        i += 1

    doc.save(output_path)
    return True


def main():
    parser = argparse.ArgumentParser(
        description="AI 聊天记录导出工具 - 从分享链接提取对话"
    )
    parser.add_argument("url", help="分享链接 (目前支持 DeepSeek)")
    parser.add_argument("-o", "--output", help="输出文件名（不含扩展名），默认从对话内容自动生成到 ./output 目录", default=None)
    parser.add_argument(
        "-f",
        "--format",
        choices=["markdown", "docx", "both"],
        default="both",
        help="输出格式: markdown, docx, 或 both (默认: both)",
    )
    parser.add_argument(
        "--headed", action="store_true", help="显示浏览器窗口 (调试用)"
    )
    parser.add_argument(
        "--timeout", type=int, default=30000, help="页面加载超时时间 (毫秒)"
    )

    args = parser.parse_args()

    # 判断平台
    url = args.url
    if "deepseek.com/share/" in url:
        platform = "deepseek"
    else:
        print(f"错误: 不支持的链接格式: {url}")
        print("目前支持: DeepSeek (chat.deepseek.com/share/...)")
        sys.exit(1)

    print(f"[信息] 正在打开链接: {url}")
    print(f"[信息] 平台: DeepSeek")
    print(f"[信息] 输出格式: {args.format}")

    with sync_playwright() as p:
        # 优先使用系统 Chrome，避免打包后找不到浏览器
        try:
            browser = p.chromium.launch(headless=not args.headed, channel="chrome")
        except Exception:
            browser = p.chromium.launch(headless=not args.headed)
        context = browser.new_context(
            viewport={"width": 1920, "height": 1080},
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36",
        )
        page = context.new_page()

        try:
            messages, meta = extract_deepseek(page, url, args.timeout)

            if not messages:
                print("[错误] 未能提取到对话内容。可能原因:")
                print("  1. 链接已过期或需要登录")
                print("  2. 页面结构已更新")
                print("  3. 网络连接问题")
                print("\n提示: 尝试使用 --headed 参数查看浏览器状态")
                sys.exit(1)

            print(f"[成功] 共提取 {len(messages)} 条消息:")
            for i, msg in enumerate(messages):
                role_cn = "用户" if msg["role"] == "user" else "AI"
                preview = msg["content"][:60].replace("\n", " ")
                print(f"  {i+1}. [{role_cn}] {preview}...")

            # 生成智能文件名
            if args.output:
                output_base = Path(args.output)
            else:
                # 确保输出到 output 目录
                output_dir = Path.cwd() / "output"
                output_dir.mkdir(exist_ok=True)

                # 按优先级选择文件名来源：
                # 1. 用户上传的文件名
                # 2. 思考过程中提取的论文标题
                # 3. 第一条用户消息
                # 4. 页面标题
                topic = ""
                
                # 优先使用文件名（去除扩展名）
                file_name = meta.get("fileName", "")
                if file_name:
                    topic = Path(file_name).stem
                
                if not topic:
                    paper_title = meta.get("paperTitle", "")
                    if paper_title:
                        topic = paper_title
                
                if not topic:
                    topic = meta.get("firstMessage", "")
                
                if not topic:
                    topic = meta.get("pageTitle", "")
                
                if not topic:
                    topic = "deepseek_chat"
                
                safe_name = sanitize_filename(topic)
                if not safe_name:
                    safe_name = "deepseek_chat"
                
                output_base = output_dir / safe_name

            print(f"[信息] 输出文件名: {output_base}")

            # 生成 Markdown
            md_content = parse_messages_to_markdown(messages)

            # 保存文件
            saved_files = []

            if args.format in ("markdown", "both"):
                md_path = output_base.with_suffix(".md")
                md_path.write_text(md_content, encoding="utf-8")
                saved_files.append(str(md_path))
                print(f"[保存] Markdown: {md_path}")

            if args.format in ("docx", "both"):
                docx_path = output_base.with_suffix(".docx")
                if markdown_to_docx(md_content, docx_path):
                    saved_files.append(str(docx_path))
                    print(f"[保存] Word: {docx_path}")
                else:
                    print("[警告] Word 文档生成失败")

            # 保存 JSON (便于后续处理)
            json_path = output_base.with_suffix(".json")
            json_path.write_text(
                json.dumps(messages, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            saved_files.append(str(json_path))
            print(f"[保存] JSON: {json_path}")

        except Exception as e:
            print(f"[错误] 提取过程中出错: {e}")
            if args.headed:
                print("浏览器窗口将保持打开，请检查页面状态")
                input("按 Enter 关闭浏览器...")
            sys.exit(1)
        finally:
            browser.close()

    print(f"\n[完成] 已保存 {len(saved_files)} 个文件")


if __name__ == "__main__":
    main()

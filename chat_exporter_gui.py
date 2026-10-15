#!/usr/bin/env python3
"""
AI 聊天记录导出工具 - 图形界面版
支持从 DeepSeek / 智谱清言 分享链接提取对话，导出为 Markdown / Word / JSON。

启动方式:
    python chat_exporter_gui.py
"""

import json
import os
import re
import subprocess
import sys
import threading
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from playwright.sync_api import sync_playwright


# ─────────────────────────── 核心逻辑 ───────────────────────────

def sanitize_filename(name, max_len=60):
    """清理文件名中的非法字符。"""
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name)
    name = re.sub(r'[_\s]+', '_', name)
    name = name.strip('_ ')
    if len(name) > max_len:
        name = name[:max_len].rstrip('_ ')
    return name


def extract_from_url(url, timeout=30000, progress_callback=None):
    """从 URL 提取对话内容，返回 (messages, meta, platform)。"""
    # 检测平台
    if "deepseek.com/share/" in url:
        platform = "deepseek"
    elif "chatglm.cn/share/" in url:
        platform = "chatglm"
    else:
        raise ValueError("不支持的链接格式，请使用 DeepSeek 或智谱清言的分享链接")

    if progress_callback:
        progress_callback("正在启动浏览器...")

    with sync_playwright() as p:
        browser = None
        try:
            try:
                browser = p.chromium.launch(headless=True, channel="chrome")
            except Exception:
                browser = p.chromium.launch(headless=True)
            context = browser.new_context(
                viewport={"width": 1920, "height": 1080},
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36",
            )
            page = context.new_page()

            if progress_callback:
                progress_callback("正在加载页面...")

            page.goto(url, wait_until="networkidle", timeout=timeout)

            if platform == "deepseek":
                try:
                    page.wait_for_selector(".ds-message", timeout=timeout)
                except Exception:
                    pass
                page.wait_for_timeout(3000)
            else:
                page.wait_for_timeout(5000)

            if progress_callback:
                progress_callback("正在提取元信息...")

            if platform == "deepseek":
                meta = page.evaluate("""() => {
                    const info = {};
                    info.pageTitle = document.title || '';
                    const firstUserMsg = document.querySelector('.ds-collapsible-text, .fbb737a4');
                    info.firstMessage = firstUserMsg ? firstUserMsg.innerText.substring(0, 150) : '';
                    const fileEl = document.querySelector('._967f3f9, .e70accd6');
                    info.fileName = fileEl ? fileEl.innerText.trim() : '';
                    const thinkEls = document.querySelectorAll('.ds-think-content');
                    let paperTitle = '';
                    for (const el of thinkEls) {
                        const text = el.innerText || '';
                        const titleMatch = text.match(/题目[：:](.+)/);
                        if (titleMatch) { paperTitle = titleMatch[1].trim(); break; }
                        const engMatch = text.match(/(Predictive design.+?)(?:[.。\\n]|$)/i);
                        if (engMatch) { paperTitle = engMatch[1].trim(); break; }
                    }
                    info.paperTitle = paperTitle;
                    return info;
                }""")
            else:
                meta = page.evaluate("""() => {
                    const info = {};
                    info.pageTitle = document.title || '';
                    const firstQuestion = document.querySelector('.question-txt');
                    info.firstMessage = firstQuestion ? firstQuestion.innerText.substring(0, 150) : '';
                    const fileEl = document.querySelector('.question-file .file-box');
                    info.fileName = fileEl ? fileEl.innerText.trim().replace(/\\s*PDF\\s*[\\d.]+MB/, '') : '';
                    return info;
                }""")

            if progress_callback:
                progress_callback("正在提取对话消息...")

            if platform == "deepseek":
                messages = page.evaluate("""() => {
                    function extractNodeText(node) {
                        if (node.nodeType === 3) {
                            return node.textContent;
                        }
                        if (node.nodeType !== 1) return '';

                        const tag = node.tagName;
                        const cls = (typeof node.className === 'string') ? node.className : '';

                        if (['SCRIPT', 'STYLE', 'BUTTON', 'SVG', 'PATH'].includes(tag)) return '';

                        if (cls.includes('katex') || cls.includes('math') || cls.includes('equation') || tag === 'MATH') {
                            const ann = node.querySelector('annotation[encoding="application/x-tex"]');
                            if (ann) return '$' + ann.textContent + '$';
                            return '$' + node.innerText.replace(/\\s+/g, ' ').trim() + '$';
                        }

                        if (tag === 'CODE' || tag === 'PRE') {
                            return '`' + node.innerText + '`';
                        }

                        if (tag === 'TABLE') {
                            return extractTable(node);
                        }

                        if (tag === 'IMG') {
                            const alt = node.alt || '';
                            const src = node.src || '';
                            return '![' + alt + '](' + src + ')';
                        }

                        if (tag === 'BR') return '\\n';

                        const blockTags = ['P', 'DIV', 'H1', 'H2', 'H3', 'H4', 'H5', 'H6',
                                           'LI', 'TR', 'BLOCKQUOTE', 'SECTION', 'ARTICLE'];
                        const isBlock = blockTags.includes(tag);

                        let result = '';
                        for (const child of node.childNodes) {
                            result += extractNodeText(child);
                        }

                        if (!cls.includes('katex') && !cls.includes('math')) {
                            result = result.replace(/  +/g, ' ');
                        }

                        if (isBlock && result.trim()) {
                            result = '\\n' + result.trim() + '\\n';
                        }

                        return result;
                    }

                    function extractTable(table) {
                        const rows = [];
                        for (const tr of table.querySelectorAll('tr')) {
                            const cells = [];
                            for (const td of tr.querySelectorAll('td, th')) {
                                cells.push(td.innerText.replace(/\\|/g, '\\\\|').trim());
                            }
                            rows.push(cells);
                        }
                        if (rows.length === 0) return '';
                        let md = '| ' + rows[0].join(' | ') + ' |\\n';
                        md += '| ' + rows[0].map(() => '---').join(' | ') + ' |\\n';
                        for (let i = 1; i < rows.length; i++) {
                            md += '| ' + rows[i].join(' | ') + ' |\\n';
                        }
                        return '\\n' + md;
                    }

                    const results = [];
                    const allMsgs = document.querySelectorAll('.ds-message');
                    for (const msg of allMsgs) {
                        const aiMain = msg.querySelector('.ds-assistant-message-main-content');
                        if (aiMain) {
                            let mainContent = extractNodeText(aiMain);
                            const thinkEl = msg.querySelector('.ds-think-content');
                            if (thinkEl) {
                                const thinkText = extractNodeText(thinkEl);
                                if (mainContent.startsWith(thinkText)) {
                                    mainContent = mainContent.substring(thinkText.length).trim();
                                }
                            }
                            mainContent = mainContent.replace(/\\n{3,}/g, '\\n\\n').trim();
                            if (mainContent) {
                                results.push({ role: 'assistant', content: mainContent });
                            }
                        } else {
                            const userEl = msg.querySelector('.ds-collapsible-text') ||
                                           msg.querySelector('.fbb737a4') ||
                                           msg.querySelector('._25c7358');
                            if (userEl) {
                                const content = userEl.innerText.trim();
                                if (content) {
                                    results.push({ role: 'user', content: content });
                                }
                            }
                        }
                    }
                    return results;
                }""")
            else:
                messages = page.evaluate("""() => {
                    const items = document.querySelectorAll('.conversation-item');
                    const messages = [];

                    for (const item of items) {
                        const questionTxt = item.querySelector('.question-txt');
                        if (questionTxt) {
                            const userText = questionTxt.innerText.trim();
                            if (userText) {
                                messages.push({role: 'user', content: userText});
                            }
                        }

                        const answerWraps = item.querySelectorAll('.answer-content-wrap:not(.text-advance-thinking-content)');
                        if (answerWraps.length > 0) {
                            const lastAnswer = answerWraps[answerWraps.length - 1];
                            const answerText = lastAnswer.innerText.trim();
                            if (answerText) {
                                messages.push({role: 'assistant', content: answerText});
                            }
                        }
                    }

                    return messages;
                }""")

        finally:
            if browser:
                try:
                    browser.close()
                except Exception:
                    pass

    return messages, meta, platform


def add_ext(base, ext):
    """为不含扩展名的路径追加扩展名。

    不使用 Path.with_suffix()，因为它会把文件名中最后一个 "." 之后的内容
    当作旧扩展名替换掉（例如 "Qwen2.5模型评测" 会变成 "Qwen2.md"）。
    若路径已带目标扩展名则原样返回。
    """
    base = Path(base)
    if base.suffix.lower() == ext:
        return base
    return base.with_name(base.name + ext)


def generate_smart_filename(meta):
    """根据元信息生成智能文件名。"""
    topic = ""
    file_name = meta.get("fileName", "")
    if file_name:
        topic = Path(file_name).stem
    if not topic:
        paper_title = meta.get("paperTitle", "")
        if paper_title:
            topic = paper_title
    if not topic:
        topic = meta.get("firstMessage", "")
    if not topic:
        topic = meta.get("pageTitle", "")
    if not topic:
        topic = "ai_chat"
    safe_name = sanitize_filename(topic)
    return safe_name if safe_name else "ai_chat"


def parse_messages_to_markdown(messages, title="AI 对话记录"):
    """将消息列表转换为 Markdown 格式，一问一答配对。
    支持将一条包含多个问题的消息拆分为多个独立问答对。
    """
    import re as _re

    def split_questions(text):
        """将包含多个编号问题的文本拆分为独立问题列表。"""
        # 匹配 1、2、3、 或 1. 2. 3. 等编号格式
        # 中文标点（、．）后不要求空格；英文句点后要求空格，避免误拆 "1.5" 之类数字
        parts = _re.split(r'\n(?=\d+[、．]\s*|\d+\.\s)', text.strip())
        if len(parts) <= 1:
            parts = _re.split(r'\n(?=(?:问题|Q)\s*\d+)', text.strip())
        return [p.strip() for p in parts if p.strip()]

    def split_answers_by_questions(text, expected_count):
        """根据问题数量拆分答案。只在答案段落数与问题数匹配时才拆分。"""
        # 中文标点（．：）后不要求空格；英文句点/冒号后要求空格
        parts = _re.split(r'\n(?=\d+[．：]\s*|\d+[.:]\s)', text.strip())
        if len(parts) <= 1:
            parts = _re.split(r'\n(?=\d+\s)', text.strip())

        # 只有当拆分出的段落数与问题数一致时才使用拆分结果
        if len(parts) == expected_count:
            return [p.strip() for p in parts if p.strip()]

        # 否则不拆分，整体作为一个答案
        return [text.strip()]

    md_lines = [f"# {title}", ""]

    i = 0
    q_num = 0
    while i < len(messages):
        msg = messages[i]
        if msg["role"] == "user":
            questions = split_questions(msg["content"])
            if i + 1 < len(messages) and messages[i + 1]["role"] == "assistant":
                answers = split_answers_by_questions(messages[i + 1]["content"], len(questions))
                # 按最大数量配对
                pairs = max(len(questions), len(answers))
                for j in range(pairs):
                    q_num += 1
                    q_text = questions[j] if j < len(questions) else "（见上文）"
                    a_text = answers[j] if j < len(answers) else "（见上文）"
                    md_lines.append(f"## 问题 {q_num}")
                    md_lines.append("")
                    md_lines.append(q_text)
                    md_lines.append("")
                    md_lines.append(f"## 回答 {q_num}")
                    md_lines.append("")
                    md_lines.append(a_text)
                    md_lines.append("")
                    md_lines.append("---")
                    md_lines.append("")
                i += 2
            else:
                for q_text in questions:
                    q_num += 1
                    md_lines.append(f"## 问题 {q_num}")
                    md_lines.append("")
                    md_lines.append(q_text)
                    md_lines.append("")
                    md_lines.append("---")
                    md_lines.append("")
                i += 1
        else:
            md_lines.append("## 回答")
            md_lines.append("")
            md_lines.append(msg["content"])
            md_lines.append("")
            md_lines.append("---")
            md_lines.append("")
            i += 1

    return "\n".join(md_lines)


def save_markdown(md_text, output_path):
    """保存 Markdown 文件。"""
    Path(output_path).write_text(md_text, encoding="utf-8")
    return str(output_path)


def save_docx(md_text, output_path):
    """保存 Word 文档。"""
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.oxml.ns import qn
    except ImportError:
        return None, "需要安装 python-docx: pip install python-docx"

    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(11)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), '宋体')

    def set_run_font(run, font_name="Times New Roman", east_asia="宋体", size=None):
        run.font.name = font_name
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = rPr.makeelement(qn('w:rFonts'), {})
            rPr.insert(0, rFonts)
        rFonts.set(qn('w:eastAsia'), east_asia)
        if size:
            run.font.size = Pt(size)

    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            heading = doc.add_heading(level=1)
            run = heading.add_run(line[2:])
            set_run_font(run, size=18)
        elif line.startswith("## "):
            heading = doc.add_heading(level=2)
            run = heading.add_run(line[3:])
            set_run_font(run, size=14)
        elif line == "---":
            doc.add_paragraph("")
        elif line.strip() == "":
            pass
        else:
            if line.strip().startswith("```"):
                code_lines = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    code_lines.append(lines[i])
                    i += 1
                if code_lines:
                    p = doc.add_paragraph()
                    run = p.add_run("\n".join(code_lines))
                    set_run_font(run, font_name="Consolas", east_asia="Consolas", size=9)
                    p.paragraph_format.left_indent = Inches(0.5)
            else:
                p = doc.add_paragraph()
                run = p.add_run(line)
                set_run_font(run)
        i += 1

    doc.save(output_path)
    return str(output_path), None


def save_json(messages, output_path):
    """保存 JSON 文件。"""
    Path(output_path).write_text(
        json.dumps(messages, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return str(output_path)


# ─────────────────────────── GUI 界面 ───────────────────────────

class ChatExporterGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("AI 聊天记录导出工具")
        self.root.geometry("720x560")
        self.root.minsize(600, 480)

        # 设置窗口图标（优先 icon.ico，其次 icon.png）
        self._set_icon()

        # 输出目录
        self.output_dir = Path.cwd() / "output"
        self.output_dir.mkdir(exist_ok=True)

        self._build_ui()

    def _set_icon(self):
        """加载窗口图标。"""
        base = Path(__file__).parent
        for name in ["icon.ico", "icon.png"]:
            icon_path = base / name
            if icon_path.exists():
                try:
                    if name.endswith(".ico"):
                        self.root.iconbitmap(str(icon_path))
                    else:
                        # tkinter 不直接支持 PNG 图标，用 PhotoImage 设置
                        img = tk.PhotoImage(file=str(icon_path))
                        self.root.iconphoto(True, img)
                        self._icon_img = img  # 防止被垃圾回收
                    return
                except Exception:
                    pass

    # ── 界面构建 ──

    def _build_ui(self):
        # 顶部区域：URL + 按钮
        frame_top = ttk.LabelFrame(self.root, text="输入分享链接", padding=10)
        frame_top.pack(fill=tk.X, padx=12, pady=(12, 6))

        self.url_var = tk.StringVar()
        entry_url = ttk.Entry(frame_top, textvariable=self.url_var, font=("Consolas", 10))
        entry_url.pack(fill=tk.X, pady=(0, 6))
        entry_url.bind("<Return>", lambda e: self._start_export())

        btn_frame = ttk.Frame(frame_top)
        btn_frame.pack(fill=tk.X)

        self.btn_export = ttk.Button(btn_frame, text="开始导出", command=self._start_export)
        self.btn_export.pack(side=tk.LEFT)
        ttk.Button(btn_frame, text="选择输出目录", command=self._choose_dir).pack(side=tk.LEFT, padx=(8, 0))

        self.dir_var = tk.StringVar(value=str(self.output_dir))
        ttk.Label(btn_frame, textvariable=self.dir_var, foreground="gray").pack(side=tk.LEFT, padx=(12, 0), fill=tk.X, expand=True)

        # 中部区域：格式选择
        frame_mid = ttk.LabelFrame(self.root, text="输出格式", padding=10)
        frame_mid.pack(fill=tk.X, padx=12, pady=6)

        self.fmt_md = tk.BooleanVar(value=True)
        self.fmt_docx = tk.BooleanVar(value=True)
        self.fmt_json = tk.BooleanVar(value=True)

        ttk.Checkbutton(frame_mid, text="Markdown (.md)", variable=self.fmt_md).pack(side=tk.LEFT, padx=(0, 16))
        ttk.Checkbutton(frame_mid, text="Word (.docx)", variable=self.fmt_docx).pack(side=tk.LEFT, padx=(0, 16))
        ttk.Checkbutton(frame_mid, text="JSON (.json)", variable=self.fmt_json).pack(side=tk.LEFT)

        # 进度条
        self.progress_var = tk.StringVar(value="就绪")
        ttk.Label(self.root, textvariable=self.progress_var, foreground="blue").pack(anchor=tk.W, padx=14, pady=(4, 0))

        self.progress_bar = ttk.Progressbar(self.root, mode="indeterminate")
        self.progress_bar.pack(fill=tk.X, padx=12, pady=4)

        # 底部区域：文件列表
        frame_bottom = ttk.LabelFrame(self.root, text="输出文件列表（双击打开）", padding=10)
        frame_bottom.pack(fill=tk.BOTH, expand=True, padx=12, pady=(6, 12))

        cols = ("文件名", "大小", "格式")
        self.tree = ttk.Treeview(frame_bottom, columns=cols, show="headings", selectmode="browse")
        self.tree.heading("文件名", text="文件名")
        self.tree.heading("大小", text="大小")
        self.tree.heading("格式", text="格式")
        self.tree.column("文件名", width=380)
        self.tree.column("大小", width=80, anchor=tk.E)
        self.tree.column("格式", width=80, anchor=tk.CENTER)
        self.tree.pack(fill=tk.BOTH, expand=True)
        self.tree.bind("<Double-1>", self._open_file)

        scrollbar = ttk.Scrollbar(frame_bottom, orient=tk.VERTICAL, command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        self._refresh_file_list()

    # ── 事件处理 ──

    def _choose_dir(self):
        d = filedialog.askdirectory(title="选择输出目录", initialdir=self.output_dir)
        if d:
            self.output_dir = Path(d)
            self.dir_var.set(str(self.output_dir))
            self._refresh_file_list()

    def _start_export(self):
        url = self.url_var.get().strip()
        if not url:
            messagebox.showwarning("提示", "请输入分享链接")
            return
        if "deepseek.com/share/" not in url and "chatglm.cn/share/" not in url:
            messagebox.showwarning("提示", "支持的链接格式:\n- DeepSeek (chat.deepseek.com/share/...)\n- 智谱清言 (chatglm.cn/share/...)")
            return

        # 在主线程中读取界面状态，避免工作线程访问 tkinter 变量（非线程安全）
        fmt_md = self.fmt_md.get()
        fmt_docx = self.fmt_docx.get()
        fmt_json = self.fmt_json.get()
        if not (fmt_md or fmt_docx or fmt_json):
            messagebox.showwarning("提示", "请至少选择一种输出格式")
            return
        output_dir = self.output_dir

        # 禁用导出按钮，防止重复点击产生并发导出
        self.btn_export.config(state=tk.DISABLED)
        self.progress_bar.start(12)
        self.progress_var.set("正在导出，请稍候...")
        threading.Thread(
            target=self._export_thread,
            args=(url, fmt_md, fmt_docx, fmt_json, output_dir),
            daemon=True,
        ).start()

    def _export_thread(self, url, fmt_md, fmt_docx, fmt_json, output_dir):
        try:
            messages, meta, platform = extract_from_url(
                url, progress_callback=lambda msg: self.root.after(0, self.progress_var.set, msg)
            )
            if not messages:
                self.root.after(0, self._on_error, "未能提取到对话内容")
                return

            safe_name = generate_smart_filename(meta)
            output_base = output_dir / safe_name

            saved = []

            # Markdown
            if fmt_md:
                md_text = parse_messages_to_markdown(messages, title=safe_name)
                path = save_markdown(md_text, add_ext(output_base, ".md"))
                saved.append(path)

            # Word
            if fmt_docx:
                md_text = parse_messages_to_markdown(messages, title=safe_name)
                path, err = save_docx(md_text, add_ext(output_base, ".docx"))
                if err:
                    self.root.after(0, messagebox.showwarning, "Word 导出失败", err)
                elif path:
                    saved.append(path)

            # JSON
            if fmt_json:
                path = save_json(messages, add_ext(output_base, ".json"))
                saved.append(path)

            self.root.after(0, self._on_done, saved)

        except Exception as e:
            self.root.after(0, self._on_error, str(e))

    def _on_done(self, saved_files):
        self.progress_bar.stop()
        self.btn_export.config(state=tk.NORMAL)
        count = len(saved_files)
        self.progress_var.set(f"导出完成，共 {count} 个文件")
        self._refresh_file_list()
        if saved_files:
            messagebox.showinfo("导出完成", f"成功导出 {count} 个文件到:\n{self.output_dir}")

    def _on_error(self, msg):
        self.progress_bar.stop()
        self.btn_export.config(state=tk.NORMAL)
        self.progress_var.set("导出失败")
        messagebox.showerror("错误", msg)

    def _refresh_file_list(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        if not self.output_dir.exists():
            return

        # 先收集 (mtime, path, size)，任何文件出错（被占用/已删除/断链）都跳过，避免界面崩溃
        entries = []
        try:
            candidates = list(self.output_dir.iterdir())
        except OSError:
            return
        for f in candidates:
            try:
                if f.suffix.lower() in (".md", ".docx", ".json"):
                    stat = f.stat()
                    entries.append((stat.st_mtime, f, stat.st_size))
            except OSError:
                continue

        for mtime, f, size in sorted(entries, key=lambda e: e[0], reverse=True):
            if size > 1024 * 1024:
                size_str = f"{size / 1024 / 1024:.1f} MB"
            elif size > 1024:
                size_str = f"{size / 1024:.1f} KB"
            else:
                size_str = f"{size} B"
            fmt = f.suffix.upper().lstrip(".")
            self.tree.insert("", tk.END, values=(f.name, size_str, fmt), iid=str(f))

    def _open_file(self, event):
        sel = self.tree.selection()
        if not sel:
            return
        file_path = sel[0]
        if not Path(file_path).exists():
            return
        # 使用系统默认程序打开文件（避免被 VS Code 拦截）
        if sys.platform == "win32":
            # 使用 ShellExecute 的 open 动词，调用系统默认关联程序
            try:
                os.startfile(file_path)
            except Exception:
                # 备用方案：通过 cmd start 命令打开
                subprocess.Popen(
                    ["cmd", "/c", "start", "", file_path],
                    shell=False,
                )
        elif sys.platform == "darwin":
            subprocess.Popen(["open", file_path])
        else:
            subprocess.Popen(["xdg-open", file_path])


# ─────────────────────────── 入口 ───────────────────────────

def main():
    root = tk.Tk()
    ChatExporterGUI(root)
    root.mainloop()


if __name__ == "__main__":
    main()
